"""SILLON - Rapport des doublons de fichiers, restreint a une taille minimale.
Variante allegee de rapport_nettoyage_disque.py : ne produit qu'un classeur
Excel (pas de PDF ni de graphique), et ne s'interesse qu'aux doublons parmi
les fichiers d'au moins TAILLE_MIN_KO, a partir du meme inventaire importe
dans SILLON (colonnes nom_fichier / hash / taille_ko / chemin_complet).

Utilite : sur un inventaire complet, la plupart des groupes de doublons sont
de petits fichiers (icones, gabarits...) sans interet pour le nettoyage -
ce script isole uniquement les doublons dont recuperer l'espace vaut la
peine (ex. tous les fichiers de 30 Mo et plus, ou seulement 100 Mo et plus).

Contrat d'execution (cahier des charges Sec 5.4, worker.py) : chaine de
connexion et repertoire de sortie fournis exclusivement par variables
d'environnement, jamais en dur.
"""
import os

import pandas as pd
import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

dsn = os.environ["SILLON_DSN"]
resultats = os.environ["SILLON_RESULTATS"]

# --- A adapter au nom reel de la table creee lors de l'import du CSV
# (visible dans l'onglet Bases > Tables) ---
NOM_TABLE = "inventaire_fichiers"

# --- Taille minimale des fichiers a analyser, en Ko (Kio, meme unite que
# la colonne taille_ko - voir rapport_nettoyage_disque.py). Exemples :
#   30 Mo  -> 30 * 1024  = 30720
#   100 Mo -> 100 * 1024 = 102400
# Seuls les fichiers d'au moins cette taille entrent dans l'analyse des
# doublons ci-dessous ; un doublon de petite taille dont l'original ou les
# copies tombent sous ce seuil est ignore.
TAILLE_MIN_KO = 30 * 1024

connexion = psycopg2.connect(dsn)
inventaire = pd.read_sql(
    f"SELECT nom_fichier, hash, taille_ko, chemin_complet FROM {NOM_TABLE}",
    connexion,
)
connexion.close()

print(f"{len(inventaire)} fichiers charges au total ({inventaire['taille_ko'].sum() / 1_048_576:.1f} Go).")

# =====================================================================
# 1. FILTRAGE PAR TAILLE MINIMALE
# =====================================================================
retenus = inventaire[inventaire["taille_ko"] >= TAILLE_MIN_KO].copy()
print(f"{len(retenus)} fichier(s) d'au moins {TAILLE_MIN_KO / 1024:.0f} Mo retenu(s) pour l'analyse.")

# =====================================================================
# 2. DOUBLONS - groupes de fichiers de contenu identique (meme hash),
#    uniquement parmi les fichiers retenus ci-dessus
# =====================================================================
comptage = retenus.groupby("hash").size()
hashs_doublons = comptage[comptage > 1].index

doublons = retenus[retenus["hash"].isin(hashs_doublons)].copy()
doublons["rang"] = doublons.groupby("hash")["chemin_complet"].rank(method="first")
doublons["statut"] = doublons["rang"].map(lambda r: "original" if r == 1 else "doublon")
doublons["taille_mo"] = (doublons["taille_ko"] / 1024).round(1)

synthese_doublons = (
    retenus[retenus["hash"].isin(hashs_doublons)]
    .groupby("hash")
    .agg(nb_copies=("nom_fichier", "count"), taille_ko=("taille_ko", "first"))
    .reset_index()
)
synthese_doublons["ko_recuperables"] = (synthese_doublons["nb_copies"] - 1) * synthese_doublons["taille_ko"]
total_ko_recuperables = synthese_doublons["ko_recuperables"].sum()

# Les groupes qui liberent le plus d'espace en premier : la liste sert a
# prioriser le nettoyage, pas juste a lister les doublons dans un ordre
# arbitraire.
rang_recuperable = synthese_doublons.set_index("hash")["ko_recuperables"]
doublons["ko_recuperables_groupe"] = doublons["hash"].map(rang_recuperable)
doublons = doublons.sort_values(
    ["ko_recuperables_groupe", "hash", "statut"], ascending=[False, True, False]
)

print(f"{len(hashs_doublons)} groupe(s) de doublons, {len(doublons) - len(hashs_doublons)} fichier(s) en trop, "
      f"{total_ko_recuperables / 1_048_576:.1f} Go recuperables.")

# =====================================================================
# 3. CLASSEUR EXCEL (2 onglets : Synthese, Doublons)
# =====================================================================
classeur = Workbook()


def ecrire_feuille(feuille, dataframe):
    feuille.append(list(dataframe.columns))
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for ligne in dataframe.itertuples(index=False):
        feuille.append(list(ligne))
    for i, colonne in enumerate(dataframe.columns, start=1):
        largeur = max(12, min(60, int(dataframe[colonne].astype(str).str.len().max()) + 2))
        feuille.column_dimensions[get_column_letter(i)].width = largeur


synthese_feuille = classeur.active
synthese_feuille.title = "Synthese"
synthese_feuille.append(["Indicateur", "Valeur"])
for cellule in synthese_feuille[1]:
    cellule.font = Font(bold=True)
synthese_feuille.append(["Taille minimale analysee (Mo)", round(TAILLE_MIN_KO / 1024, 1)])
synthese_feuille.append(["Fichiers analyses (>= seuil)", len(retenus)])
synthese_feuille.append(["Volume analyse (Go)", round(retenus["taille_ko"].sum() / 1_048_576, 2)])
synthese_feuille.append(["Groupes de doublons", len(hashs_doublons)])
synthese_feuille.append(["Fichiers en double", len(doublons) - len(hashs_doublons)])
synthese_feuille.append(["Espace recuperable (Go)", round(total_ko_recuperables / 1_048_576, 2)])

feuille_doublons = classeur.create_sheet("Doublons")
ecrire_feuille(
    feuille_doublons,
    doublons[["hash", "statut", "nom_fichier", "chemin_complet", "taille_mo", "taille_ko"]],
)

classeur.save(os.path.join(resultats, "rapport_doublons_taille_minimale.xlsx"))

print("\nFichier produit : rapport_doublons_taille_minimale.xlsx")
