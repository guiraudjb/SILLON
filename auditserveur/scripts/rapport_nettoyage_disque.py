"""SILLON - Rapport d'aide au nettoyage d'espace disque.
Doublons de fichiers (par hash), taille cumulee par dossier/sous-dossier,
fichiers les plus volumineux. Genere un PDF de synthese et un classeur
Excel de detail, a partir d'un inventaire de fichiers importe dans SILLON
(colonnes nom_fichier / hash / taille_ko / chemin_complet).

Contrat d'execution (cahier des charges Sec 5.4, worker.py) : chaine de
connexion et repertoire de sortie fournis exclusivement par variables
d'environnement, jamais en dur.
"""
import os
from collections import defaultdict

import matplotlib
matplotlib.use("Agg")  # aucun affichage interactif possible dans le conteneur d'execution
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import pandas as pd
import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

dsn = os.environ["SILLON_DSN"]
resultats = os.environ["SILLON_RESULTATS"]

# --- A adapter au nom reel de la table creee lors de l'import du CSV
# (visible dans l'onglet Bases > Tables) ---
NOM_TABLE = "inventaire_fichiers"

connexion = psycopg2.connect(dsn)
inventaire = pd.read_sql(
    f"SELECT nom_fichier, hash, taille_ko, chemin_complet FROM {NOM_TABLE}",
    connexion,
)
connexion.close()

print(f"{len(inventaire)} fichiers charges, {inventaire['taille_ko'].sum() / 1_048_576:.1f} Go au total.")

# =====================================================================
# 1. DOUBLONS - groupes de fichiers de contenu identique (meme hash)
# =====================================================================
comptage = inventaire.groupby("hash").size()
hashs_doublons = comptage[comptage > 1].index

doublons = inventaire[inventaire["hash"].isin(hashs_doublons)].copy()
doublons["rang"] = doublons.groupby("hash")["chemin_complet"].rank(method="first")
doublons["statut"] = doublons["rang"].map(lambda r: "original" if r == 1 else "doublon")

synthese_doublons = (
    inventaire[inventaire["hash"].isin(hashs_doublons)]
    .groupby("hash")
    .agg(nb_copies=("nom_fichier", "count"), taille_ko=("taille_ko", "first"))
    .reset_index()
)
synthese_doublons["ko_recuperables"] = (synthese_doublons["nb_copies"] - 1) * synthese_doublons["taille_ko"]
synthese_doublons = synthese_doublons.sort_values("ko_recuperables", ascending=False)

total_ko_recuperables = synthese_doublons["ko_recuperables"].sum()
print(f"{len(hashs_doublons)} groupe(s) de doublons, {len(doublons) - len(hashs_doublons)} fichier(s) en trop, "
      f"{total_ko_recuperables / 1_048_576:.1f} Go recuperables.")

# =====================================================================
# 2. TAILLE PAR DOSSIER ET SOUS-DOSSIER (cumulatif, tous niveaux)
# =====================================================================
def chemins_ancetres(chemin_fichier):
    segments = [s for s in chemin_fichier.split("/") if s]
    return ["/" + "/".join(segments[:i]) for i in range(1, len(segments))]


taille_dossiers = defaultdict(int)
fichiers_dossiers = defaultdict(int)
for ligne in inventaire.itertuples():
    for dossier in chemins_ancetres(ligne.chemin_complet):
        taille_dossiers[dossier] += ligne.taille_ko
        fichiers_dossiers[dossier] += 1

dossiers = pd.DataFrame(
    {"dossier": list(taille_dossiers.keys()), "taille_ko": list(taille_dossiers.values())}
)
dossiers["nb_fichiers"] = dossiers["dossier"].map(fichiers_dossiers)
dossiers["taille_go"] = (dossiers["taille_ko"] / 1_048_576).round(2)
dossiers = dossiers.sort_values("taille_ko", ascending=False).reset_index(drop=True)

# =====================================================================
# 3. FICHIERS LES PLUS VOLUMINEUX
# =====================================================================
plus_volumineux = inventaire.sort_values("taille_ko", ascending=False).head(50).copy()
plus_volumineux["taille_mo"] = (plus_volumineux["taille_ko"] / 1024).round(1)

# =====================================================================
# 4. GRAPHIQUES + RAPPORT PDF (synthese pour le service decisionnaire)
# =====================================================================
top_dossiers = dossiers.head(15).iloc[::-1]
top_doublons = synthese_doublons.head(15).iloc[::-1]
top_fichiers = plus_volumineux.head(15).iloc[::-1]

figure_synthese, axes = plt.subplots(2, 2, figsize=(12, 10))

axes[0, 0].barh(top_dossiers["dossier"], top_dossiers["taille_go"], color="#0055a4")
axes[0, 0].set_title("Top 15 dossiers par taille cumulee")
axes[0, 0].set_xlabel("Go")

axes[0, 1].barh(top_doublons["hash"].str.slice(0, 10) + "...", top_doublons["ko_recuperables"] / 1_048_576, color="#c9191e")
axes[0, 1].set_title("Top 15 groupes de doublons (espace recuperable)")
axes[0, 1].set_xlabel("Go recuperables")

axes[1, 0].barh(top_fichiers["nom_fichier"], top_fichiers["taille_mo"], color="#00915a")
axes[1, 0].set_title("Top 15 fichiers les plus volumineux")
axes[1, 0].set_xlabel("Mo")

axes[1, 1].axis("off")
resume_texte = (
    f"Fichiers analyses : {len(inventaire):,}\n"
    f"Volume total : {inventaire['taille_ko'].sum() / 1_048_576:.1f} Go\n\n"
    f"Groupes de doublons : {len(hashs_doublons)}\n"
    f"Fichiers en double : {len(doublons) - len(hashs_doublons)}\n"
    f"Espace recuperable : {total_ko_recuperables / 1_048_576:.1f} Go\n\n"
    f"Dossiers repertories : {len(dossiers)}\n"
).replace(",", " ")
axes[1, 1].text(0.05, 0.95, resume_texte, fontsize=12, va="top", family="monospace")
axes[1, 1].set_title("Resume")

figure_synthese.suptitle("Rapport d'aide au nettoyage d'espace disque", fontsize=14)
figure_synthese.tight_layout()

with PdfPages(os.path.join(resultats, "rapport_nettoyage_disque.pdf")) as pdf:
    pdf.savefig(figure_synthese)
figure_synthese.savefig(os.path.join(resultats, "rapport_nettoyage_disque.png"), dpi=110, bbox_inches="tight")

# =====================================================================
# 5. CLASSEUR EXCEL DE DETAIL (pour exploitation par le service decisionnaire)
# =====================================================================
classeur = Workbook()


def ecrire_feuille(feuille, dataframe):
    feuille.append(list(dataframe.columns))
    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
    for ligne in dataframe.itertuples(index=False):
        feuille.append(list(ligne))
    for i, colonne in enumerate(dataframe.columns, start=1):
        largeur = max(12, min(60, int(dataframe[colonne].astype(str).str.len().max()) + 2))
        feuille.column_dimensions[get_column_letter(i)].width = largeur


synthese_feuille = classeur.active
synthese_feuille.title = "Synthese"
synthese_feuille.append(["Indicateur", "Valeur"])
for cellule in synthese_feuille[1]:
    cellule.font = Font(bold=True)
synthese_feuille.append(["Fichiers analyses", len(inventaire)])
synthese_feuille.append(["Volume total (Go)", round(inventaire["taille_ko"].sum() / 1_048_576, 2)])
synthese_feuille.append(["Groupes de doublons", len(hashs_doublons)])
synthese_feuille.append(["Fichiers en double", len(doublons) - len(hashs_doublons)])
synthese_feuille.append(["Espace recuperable (Go)", round(total_ko_recuperables / 1_048_576, 2)])

feuille_doublons = classeur.create_sheet("Doublons")
ecrire_feuille(
    feuille_doublons,
    doublons[["hash", "statut", "nom_fichier", "chemin_complet", "taille_ko"]].sort_values(["hash", "statut"]),
)

feuille_dossiers = classeur.create_sheet("Dossiers")
ecrire_feuille(feuille_dossiers, dossiers[["dossier", "nb_fichiers", "taille_ko", "taille_go"]])

feuille_volumineux = classeur.create_sheet("Plus volumineux")
ecrire_feuille(feuille_volumineux, plus_volumineux[["nom_fichier", "chemin_complet", "taille_mo", "taille_ko"]])

classeur.save(os.path.join(resultats, "rapport_nettoyage_disque.xlsx"))

print("\nFichiers produits : rapport_nettoyage_disque.pdf, rapport_nettoyage_disque.png, "
      "rapport_nettoyage_disque.xlsx")
